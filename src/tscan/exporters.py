from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def flatten_json(value: Any, prefix: str = "$") -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            rows.extend(flatten_json(child, f"{prefix}.{key}"))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            rows.extend(flatten_json(child, f"{prefix}[{index}]"))
        if not value:
            rows.append((prefix, "[]"))
    else:
        rows.append((prefix, value))
    return rows


def export_csv(value: Any, path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["path", "value"])
        for key, child in flatten_json(value):
            writer.writerow([key, "" if child is None else child])


def export_xlsx(value: Any, path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extraction"
    sheet.append(["Path", "Value"])
    header = sheet[1]
    for cell in header:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3B64")
    for key, child in flatten_json(value):
        if isinstance(child, (dict, list)):
            child = json.dumps(child, ensure_ascii=False)
        sheet.append([key, child])
    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 64
    sheet.freeze_panes = "A2"
    workbook.save(path)
